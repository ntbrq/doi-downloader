import csv
from pathlib import Path

from openpyxl import load_workbook


def parse_dois(file_path: Path) -> list[str]:
    """从文件解析 DOI 列表，支持 txt/csv/xlsx，返回去重列表。"""
    suffix = file_path.suffix.lower()
    if suffix == ".txt":
        return _parse_txt(file_path)
    elif suffix == ".csv":
        return _parse_csv(file_path)
    elif suffix == ".xlsx":
        return _parse_xlsx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")


def _parse_txt(file_path: Path) -> list[str]:
    dois: list[str] = []
    seen: set[str] = set()
    for line in file_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or line.startswith("//"):
            continue
        if line not in seen:
            seen.add(line)
            dois.append(line)
    return dois


def _parse_csv(file_path: Path) -> list[str]:
    dois: list[str] = []
    seen: set[str] = set()
    with file_path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None or "doi" not in [h.lower() for h in reader.fieldnames]:
            raise ValueError(f"CSV file must have a 'doi' column. Found: {reader.fieldnames}")
        doi_field = next(h for h in reader.fieldnames if h.lower() == "doi")
        for row in reader:
            doi = row[doi_field].strip()
            if doi and doi not in seen:
                seen.add(doi)
                dois.append(doi)
    return dois


def _parse_xlsx(file_path: Path) -> list[str]:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    if "doi" not in header:
        raise ValueError(f"Excel file must have a 'doi' column. Found: {header}")
    doi_idx = header.index("doi")
    dois: list[str] = []
    seen: set[str] = set()
    for row in rows[1:]:
        if doi_idx >= len(row) or row[doi_idx] is None:
            continue
        doi = str(row[doi_idx]).strip()
        if doi and doi not in seen:
            seen.add(doi)
            dois.append(doi)
    wb.close()
    return dois
